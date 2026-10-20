# -*- coding: utf-8 -*-
"""
fig1_make_editable_pptx_20260916.py —— Fig. 1 的 PowerPoint 可编辑版

做法（与 Fig. 1 同源）：幻灯片尺寸 = 最终印刷尺寸（190 × 128.3 mm），
  * 15 张脑切片 + 色标条 → 600 dpi 位图（数据图无法用原生形状表达）
  * 全部文字（z 层面标注、行标签、面板字母、色标刻度与标题）→ **原生文本框**，可改字改位
  * 行标签的 n 用独立 run 设斜体

几何全部取自 fig1_lesion_distribution_20260916.py 的常量，保证与 PDF 严格同源。
"""
import os
import sys
import tempfile

import nibabel as nib
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Mm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import matplotlib as mpl
import matplotlib.pyplot as plt

import fig1_lesion_distribution_20260916 as F2
import figstyle_20260916 as fs

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Fig1_lesion_distribution_editable.pptx')
TMP = tempfile.mkdtemp(prefix='fig2pptx_')

FIG_W, FIG_H = F2.FIG_W_MM, None      # FIG_H 由实际版面算出


def build_assets():
    """导出 15 张切片位图 + 色标条位图，并返回几何参数"""
    summ = list(__import__('csv').DictReader(
        open(os.path.join(HERE, 'fig1_tertile_summary.csv'), encoding='utf-8-sig')))
    probs, aff = [], None
    for s in summ:
        img = nib.load(os.path.join(HERE, s['out_nii']))
        probs.append(np.asarray(img.dataobj, dtype=np.float32))
        aff = np.asarray(img.affine)
    from nibabel.processing import resample_from_to
    tpl = np.clip(np.asarray(resample_from_to(nib.load(F2.TEMPLATE), (probs[0].shape, aff),
                                              order=1).dataobj, dtype=np.float32), 0, 255) / 255.0
    vmax = max(float(p.max()) for p in probs)
    d0 = F2.crop_slice(probs[0], aff, 0)
    slice_h = F2.SLICE_W_MM * d0.shape[0] / float(d0.shape[1])
    fig_h = F2.GRID_BOTTOM_MM + 2 * F2.ROW_PITCH_MM + slice_h + 6.0

    # 2026-10-08 修正：原先此处**硬编码**旧 ColorBrewer 蓝，与印刷版（暖臂）不一致，
    # 且使 PPTX 完全吃不到 figstyle 的配色变更。改为直接取契约变量。
    cmap_ov = fs.CMAP_OVERLAY.copy()

    paths = {}
    old = mpl.rcParams['savefig.bbox']
    mpl.rcParams['savefig.bbox'] = None
    for r in range(3):
        for c, z in enumerate(F2.SLICE_Z):
            fig = plt.figure(figsize=(F2.SLICE_W_MM * fs.MM, slice_h * fs.MM))
            ax = fig.add_axes([0, 0, 1, 1])
            t = F2.crop_slice(tpl, aff, z)
            p = F2.crop_slice(probs[r], aff, z)
            ax.imshow(np.ma.masked_where(t <= 0.11, t), cmap='gray', vmin=0.06, vmax=0.80,
                      alpha=0.60, origin='lower', interpolation='nearest', aspect='auto')
            rgba = cmap_ov(np.clip(p / vmax, 0, 1))
            rgba[..., 3] = np.where(p > 0, 0.75, 0.0)
            ax.imshow(rgba, origin='lower', interpolation='nearest', aspect='auto')
            ax.set_axis_off()
            f = os.path.join(TMP, 's_%d_%d.png' % (r, c))
            fig.savefig(f, dpi=600)
            plt.close(fig)
            paths[(r, c)] = f
    # 色标条
    fig = plt.figure(figsize=(F2.CBAR_W_MM * fs.MM, F2.CBAR_H_MM * fs.MM))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.imshow(np.linspace(0, 1, 512).reshape(1, -1), cmap=cmap_ov, aspect='auto',
              origin='lower')
    ax.set_axis_off()
    cb_path = os.path.join(TMP, 'cbar.png')
    fig.savefig(cb_path, dpi=600)
    plt.close(fig)
    mpl.rcParams['savefig.bbox'] = old
    return paths, cb_path, slice_h, fig_h, vmax


def tb(slide, x_mm, y_mm, w_mm, h_mm):
    box = slide.shapes.add_textbox(Mm(x_mm), Mm(y_mm), Mm(w_mm), Mm(h_mm))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return box, tf


def set_run(run, text, size=7, bold=False, italic=False, color=(0x22, 0x22, 0x22)):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(*color)


def main():
    paths, cb_path, slice_h, fig_h, vmax = build_assets()
    print('  切片位图 %d 张；版面 %.1f × %.1f mm' % (len(paths), FIG_W, fig_h))

    prs = Presentation()
    prs.core_properties.last_modified_by = ''   # 2026-10-08：清 python-pptx 默认模板带的
    prs.core_properties.author = ''              #   第三方人名（Steve Canny）
    prs.slide_width = Mm(FIG_W)
    prs.slide_height = Mm(fig_h)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    labels = [('Small', '< 1.2 mL', 41), ('Medium', '1.2–6.5 mL', 42), ('Large', '> 6.5 mL', 41)]
    for r in range(3):
        y0 = F2.GRID_BOTTOM_MM + (2 - r) * F2.ROW_PITCH_MM
        ppt_y = fig_h - y0 - slice_h          # pptx 原点在左上
        for c, z in enumerate(F2.SLICE_Z):
            x0 = F2.LABEL_COL_MM + c * F2.SLICE_W_MM
            slide.shapes.add_picture(paths[(r, c)], Mm(x0), Mm(ppt_y),
                                     Mm(F2.SLICE_W_MM), Mm(slice_h))
            if r == 0:
                _, tf = tb(slide, x0, ppt_y - 4.6, F2.SLICE_W_MM, 4.0)
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                set_run(p.add_run(), 'z = %+d' % z if z else 'z = 0', size=6.8)
        # 行标签（两段：名称 + 斜体 n）
        _, tf = tb(slide, 0.5, ppt_y + slice_h / 2 - 4.0, F2.LABEL_COL_MM - 2.0, 8.0)
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.RIGHT
        set_run(p1.add_run(), '%s  (%s)' % (labels[r][0], labels[r][1]), size=7)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        set_run(p2.add_run(), 'n', size=7, italic=True)
        set_run(p2.add_run(), ' = %d' % labels[r][2], size=7)
        # 面板字母
        _, tfa = tb(slide, 0.5, ppt_y - 5.4, F2.LABEL_COL_MM - 2.0, 4.5)
        pa = tfa.paragraphs[0]
        pa.alignment = PP_ALIGN.RIGHT
        set_run(pa.add_run(), 'abc'[r], size=8, bold=True)

    # 色标
    cb_ppt_y = fig_h - F2.CBAR_Y_MM - F2.CBAR_H_MM
    slide.shapes.add_picture(cb_path, Mm(F2.CBAR_X_MM), Mm(cb_ppt_y),
                             Mm(F2.CBAR_W_MM), Mm(F2.CBAR_H_MM))
    _, tf = tb(slide, F2.CBAR_X_MM, cb_ppt_y - 4.6, 60, 4.0)
    set_run(tf.paragraphs[0].add_run(), 'Lesion overlap probability', size=6.8)
    for frac in (0.0, 0.05, 0.10, 0.15, 0.20):
        xx = F2.CBAR_X_MM + F2.CBAR_W_MM * frac / vmax
        _, tft = tb(slide, xx - 5.0, cb_ppt_y + F2.CBAR_H_MM + 0.8, 10.0, 4.0)
        pt = tft.paragraphs[0]
        pt.alignment = PP_ALIGN.CENTER
        set_run(pt.add_run(), ('%.2f' % frac).rstrip('0').rstrip('.'), size=6.3)

    prs.save(OUT)
    print('  已写出 %s' % OUT)
    shapes = len(prs.slides[0].shapes)
    pics = sum(1 for s in prs.slides[0].shapes if s.shape_type == 13)
    txts = sum(1 for s in prs.slides[0].shapes if s.has_text_frame)
    print('  形状总数 %d（图片 %d = 15 切片 + 1 色标；文本框 %d）' % (shapes, pics, txts))
    assert pics == 16, pics
    assert txts == 5 + 3 + 3 + 1 + 5, txts
    print('  结构自检 PASS')
    print('  幻灯片尺寸 %.1f × %.1f mm（= 印刷尺寸，导出无需缩放）' % (FIG_W, fig_h))


if __name__ == '__main__':
    main()
